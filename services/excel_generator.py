import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
import datetime

def generate_excel(data_list, output_path, stats=None):
    """
    Genera el archivo Excel con formato profesional, incluyendo los textos
    originales y su respectiva traducción al español si el documento está en inglés.
    """
    wb = openpyxl.Workbook()
    
    # 1. Hoja Principal
    ws = wb.active
    ws.title = "Artículos"
    
    headers = [
        "Artículo", 
        "Idioma", 
        "Nombre (Original)", 
        "Nombre (Español)", 
        "Autor", 
        "Año", 
        "Resumen (Original)", 
        "Resumen (Español)", 
        "Punto clave más relevantes (Original)", 
        "Puntos clave (Español)", 
        "Referencias Excluidas"
    ]
    ws.append(headers)
    
    # Estilos
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid") # Azul elegante
    header_font = Font(color="FFFFFF", bold=True, size=11)
    
    thin_border = Border(left=Side(style='thin', color='D9D9D9'), 
                         right=Side(style='thin', color='D9D9D9'), 
                         top=Side(style='thin', color='D9D9D9'), 
                         bottom=Side(style='thin', color='D9D9D9'))
                         
    wrap_alignment = Alignment(wrap_text=True, vertical="top")
    center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # Aplicar formato a encabezados
    for col_num, cell in enumerate(ws[1], 1):
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_alignment
        cell.border = thin_border
    
    ws.row_dimensions[1].height = 32
    
    # Agregar datos
    for row_idx, data in enumerate(data_list, start=2):
        refs_status = "Sí (Excluidas)" if data.get("Referencias_Excluidas") else "No detectadas"
        idioma_raw = data.get("Idioma", "Español")
        idioma_label = "Inglés (Traducido al español)" if idioma_raw == "Inglés" else "Español"
        
        row = [
            data.get("Artículo", ""),
            idioma_label,
            data.get("Nombre", ""),
            data.get("Nombre_ES", data.get("Nombre", "")),
            data.get("Autor", ""),
            data.get("Año", ""),
            data.get("Resumen", ""),
            data.get("Resumen_ES", data.get("Resumen", "")),
            data.get("Punto clave más relevantes", ""),
            data.get("Puntos_Clave_ES", data.get("Punto clave más relevantes", "")),
            refs_status
        ]
        ws.append(row)
        
        # Formato de celdas de datos
        for col_num in range(1, 12):
            cell = ws.cell(row=row_idx, column=col_num)
            cell.alignment = center_alignment if col_num in (1, 2, 6, 11) else wrap_alignment
            cell.border = thin_border
            cell.font = Font(size=10, name="Calibri")
            
            # Alternar color de fila
            if row_idx % 2 == 0:
                cell.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    
    # Ajustar ancho de columnas
    column_widths = {
        'A': 18, # Artículo
        'B': 22, # Idioma
        'C': 40, # Nombre Original
        'D': 40, # Nombre Español
        'E': 30, # Autor
        'F': 12, # Año
        'G': 65, # Resumen Original
        'H': 65, # Resumen Español
        'I': 70, # Puntos clave Original
        'J': 70, # Puntos clave Español
        'K': 20  # Referencias Excluidas
    }
    
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width
        
    # Congelar panel (primera fila)
    ws.freeze_panes = "A2"
    
    # Crear Tabla con Autofiltro si hay datos
    if data_list:
        max_row = len(data_list) + 1
        tab = Table(displayName="TablaArticulos", ref=f"A1:K{max_row}")
        style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                               showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        tab.tableStyleInfo = style
        ws.add_table(tab)
    
    # 2. Hoja de Información
    ws_info = wb.create_sheet(title="Información")
    ws_info.append(["Propiedad", "Valor"])
    
    info_data = [
        ["Fecha de procesamiento", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        ["Versión de la aplicación", "2.1.0 (Buscador Semántico & Traducción)"],
        ["Filtro de referencias", "Exclusión activa de bibliografía en búsqueda"],
        ["Soporte de idiomas", "Detección de idioma y traducción automática a español"]
    ]
    
    if stats:
        info_data.extend([
            ["Número total de artículos", stats.get("total", 0)],
            ["Artículos procesados", stats.get("procesados", 0)],
            ["Artículos con errores", stats.get("errores", 0)],
            ["Modelo de IA / Motor utilizado", stats.get("modelo", "Semántico Local")],
            ["Consulta semántica realizada", stats.get("consulta", "Ninguna")]
        ])
        
    for row in info_data:
        ws_info.append(row)
        
    for cell in ws_info[1]:
        cell.font = Font(bold=True)
    ws_info.column_dimensions['A'].width = 30
    ws_info.column_dimensions['B'].width = 50
    
    wb.save(output_path)
    return output_path
