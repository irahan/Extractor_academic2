import docx
import logging

logger = logging.getLogger(__name__)

def extract_text_from_docx(file_path):
    """
    Extrae el texto de un archivo DOCX.
    """
    text = ""
    try:
        doc = docx.Document(file_path)
        metadata = {
            "title": doc.core_properties.title if doc.core_properties.title else "",
            "author": doc.core_properties.author if doc.core_properties.author else ""
        }
        
        # Extraer texto de párrafos
        for para in doc.paragraphs:
            para_text = para.text.strip()
            if para_text:
                text += para_text + "\n\n"
                
        # Extraer texto de tablas (opcional pero útil)
        for table in doc.tables:
            for row in table.rows:
                row_data = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_data.append(cell.text.strip())
                if row_data:
                    text += " | ".join(row_data) + "\n"
            text += "\n"
            
        if not text.strip():
            logger.warning(f"El archivo {file_path} está vacío o no contiene texto.")
            return None, metadata
            
        return text, metadata
    except Exception as e:
        logger.error(f"Error al leer DOCX {file_path}: {e}")
        raise
